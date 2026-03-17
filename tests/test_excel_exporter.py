"""
Tests for exporters/excel_exporter.py.

Covers:
  - Sheet names are sanitized (illegal chars replaced, 31-char truncation).
  - Reports grouped by workbook_groups appear in the same file.
  - Ungrouped reports each get their own file.
  - {date} in filename is replaced with YYYY-MM-DD.
  - Failed report produces a placeholder sheet (not an exception).
  - ExportedFile.report_names contains all report names in the file.
  - Empty DataFrame produces "(no data)" sheet.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from report_runner.config.models import WorkbookGroup
from report_runner.exporters.excel_exporter import ExcelExporter, _sanitize_sheet_name
from report_runner.models import ReportResult


# ── Sheet name sanitization ───────────────────────────────────────────────────

@pytest.mark.parametrize("name,expected", [
    ("Normal Name", "Normal Name"),
    ("Name:With:Colons", "Name_With_Colons"),
    ("Name/With/Slashes", "Name_With_Slashes"),
    ("Name\\Backslash", "Name_Backslash"),
    ("Name?Question", "Name_Question"),
    ("Name*Asterisk", "Name_Asterisk"),
    ("Name[Bracket]", "Name_Bracket_"),
    # Exactly 31 chars: unchanged.
    ("A" * 31, "A" * 31),
    # 35 chars: truncated to 31.
    ("A" * 35, "A" * 31),
])
def test_sanitize_sheet_name(name: str, expected: str) -> None:
    assert _sanitize_sheet_name(name) == expected


# ── Helpers ───────────────────────────────────────────────────────────────────

def make_result(
    name: str,
    df: pd.DataFrame | None = None,
    status: str = "success",
) -> ReportResult:
    now = datetime.now()
    return ReportResult(
        report_name=name,
        status=status,  # type: ignore[arg-type]
        dataframe=df if df is not None else pd.DataFrame({"col": [1, 2]}),
        row_count=len(df) if df is not None else 2,
        started_at=now,
        finished_at=now,
    )


# ── Grouped workbooks ─────────────────────────────────────────────────────────

def test_grouped_reports_in_same_file(tmp_path: Path) -> None:
    """Reports in the same workbook_group appear as sheets in one file."""
    results = [
        make_result("Report A"),
        make_result("Report B"),
    ]
    groups = [
        WorkbookGroup(workbook="combined.xlsx", reports=["Report A", "Report B"])
    ]

    exporter = ExcelExporter()
    exported = exporter.export(groups=groups, results=results, output_dir=str(tmp_path))

    assert len(exported) == 1
    wb = load_workbook(exported[0].path)
    assert "Report A" in wb.sheetnames
    assert "Report B" in wb.sheetnames


def test_exported_file_contains_correct_report_names(tmp_path: Path) -> None:
    """ExportedFile.report_names lists all reports in the file."""
    results = [make_result("Alpha"), make_result("Beta")]
    groups = [WorkbookGroup(workbook="out.xlsx", reports=["Alpha", "Beta"])]

    exporter = ExcelExporter()
    exported = exporter.export(groups=groups, results=results, output_dir=str(tmp_path))

    assert set(exported[0].report_names) == {"Alpha", "Beta"}


# ── Ungrouped reports ─────────────────────────────────────────────────────────

def test_ungrouped_report_gets_own_file(tmp_path: Path) -> None:
    """A report not in any group gets its own .xlsx file."""
    results = [make_result("Orphan Report")]

    exporter = ExcelExporter()
    exported = exporter.export(groups=[], results=results, output_dir=str(tmp_path))

    assert len(exported) == 1
    assert exported[0].report_names == ["Orphan Report"]


def test_mixed_grouped_and_ungrouped(tmp_path: Path) -> None:
    """Grouped reports share a file; ungrouped reports each get their own."""
    results = [
        make_result("Grouped A"),
        make_result("Grouped B"),
        make_result("Ungrouped C"),
    ]
    groups = [WorkbookGroup(workbook="grouped.xlsx", reports=["Grouped A", "Grouped B"])]

    exporter = ExcelExporter()
    exported = exporter.export(groups=groups, results=results, output_dir=str(tmp_path))

    assert len(exported) == 2
    names_in_grouped = exported[0].report_names if len(exported[0].report_names) == 2 else exported[1].report_names
    assert set(names_in_grouped) == {"Grouped A", "Grouped B"}


# ── Date placeholder ──────────────────────────────────────────────────────────

def test_date_placeholder_replaced_in_filename(tmp_path: Path) -> None:
    """The {date} placeholder in workbook name is replaced with YYYY-MM-DD."""
    results = [make_result("Report")]
    groups = [WorkbookGroup(workbook="report_{date}.xlsx", reports=["Report"])]

    exporter = ExcelExporter()
    exported = exporter.export(groups=groups, results=results, output_dir=str(tmp_path))

    import re
    assert re.search(r"report_\d{4}-\d{2}-\d{2}\.xlsx", exported[0].path.name)


# ── Failed report placeholder ─────────────────────────────────────────────────

def test_failed_report_produces_placeholder_sheet(tmp_path: Path) -> None:
    """A failed report generates a sheet with an error message, not empty content."""
    results = [
        ReportResult(
            report_name="Failed Report",
            status="failed",
            dataframe=None,
            row_count=0,
            started_at=datetime.now(),
            finished_at=datetime.now(),
            error_message="Something went wrong",
        )
    ]

    exporter = ExcelExporter()
    exported = exporter.export(groups=[], results=results, output_dir=str(tmp_path))

    wb = load_workbook(exported[0].path)
    ws = wb.active
    cell_value = ws.cell(1, 1).value
    assert cell_value is not None
    assert "failed" in cell_value.lower()


# ── Empty DataFrame ───────────────────────────────────────────────────────────

def test_empty_dataframe_produces_no_data_cell(tmp_path: Path) -> None:
    """An empty DataFrame result produces a sheet with '(no data)'."""
    empty_df = pd.DataFrame()
    results = [
        ReportResult(
            report_name="Empty",
            status="success",
            dataframe=empty_df,
            row_count=0,
            started_at=datetime.now(),
            finished_at=datetime.now(),
        )
    ]

    exporter = ExcelExporter()
    exported = exporter.export(groups=[], results=results, output_dir=str(tmp_path))

    wb = load_workbook(exported[0].path)
    ws = wb.active
    assert ws.cell(1, 1).value == "(no data)"


# ── Header styling ────────────────────────────────────────────────────────────

def test_header_row_is_bold(tmp_path: Path) -> None:
    """The header row cells have bold font."""
    results = [make_result("Styled", df=pd.DataFrame({"x": [1], "y": [2]}))]

    exporter = ExcelExporter()
    exported = exporter.export(groups=[], results=results, output_dir=str(tmp_path))

    wb = load_workbook(exported[0].path)
    ws = wb.active
    assert ws.cell(1, 1).font.bold is True
    assert ws.cell(1, 2).font.bold is True
