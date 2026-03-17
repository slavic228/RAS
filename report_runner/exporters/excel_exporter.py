"""
ExcelExporter — writes report DataFrames to .xlsx files using openpyxl.

Design decisions:
  - Uses openpyxl directly (no pandas.to_excel()) for full control over styling.
  - Header row: bold, blue fill (#D9E1F2), auto-filter, frozen pane.
  - Column widths: auto-fitted to the widest cell in each column.
  - Sheet names: report name truncated to 31 chars, illegal chars → '_'.
  - Filename {date} placeholder: replaced with YYYY-MM-DD of run date.
  - Reports not assigned to any workbook group get their own file automatically.
  - Failed reports: placeholder sheet with a single descriptive cell.

The exporter is entirely decoupled from email/distribution logic.
It writes files and returns ExportedFile metadata — nothing more.
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import pandas as pd
from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from report_runner.config.models import WorkbookGroup
from report_runner.exporters.base import Exporter
from report_runner.models import ExportedFile, ReportResult

# Header styling constants
_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9E1F2",
)

# Characters illegal in Excel sheet names
_ILLEGAL_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")
_MAX_SHEET_NAME_LENGTH = 31


class ExcelExporter(Exporter):
    """Writes DataFrames to styled .xlsx files."""

    def export(
        self,
        groups: list[WorkbookGroup],
        results: list[ReportResult],
        output_dir: str,
    ) -> list[ExportedFile]:
        run_date = date.today()
        out_dir = Path(output_dir)
        out_dir.mkdir(parents=True, exist_ok=True)

        # Build a lookup so we can quickly find results by report name.
        results_by_name: dict[str, ReportResult] = {r.report_name: r for r in results}

        exported_files: list[ExportedFile] = []

        # ── Step 1: Write explicitly grouped workbooks ────────────────────────
        assigned_report_names: set[str] = set()

        for group in groups:
            file_path = self._resolve_path(out_dir, group.workbook, run_date)
            wb_results = [
                results_by_name[name]
                for name in group.reports
                if name in results_by_name
            ]
            if not wb_results:
                logger.warning(
                    "Workbook group '{}' has no matching results — skipping.",
                    group.workbook,
                )
                continue

            self._write_workbook(file_path, wb_results)
            exported_files.append(
                ExportedFile(path=file_path, report_names=[r.report_name for r in wb_results])
            )
            assigned_report_names.update(r.report_name for r in wb_results)
            logger.info("Exported workbook: {}", file_path)

        # ── Step 2: Each ungrouped report gets its own file ───────────────────
        for result in results:
            if result.report_name in assigned_report_names:
                continue

            safe_name = _sanitize_sheet_name(result.report_name)
            filename = f"{safe_name}_{run_date.strftime('%Y-%m-%d')}.xlsx"
            file_path = out_dir / filename

            self._write_workbook(file_path, [result])
            exported_files.append(
                ExportedFile(path=file_path, report_names=[result.report_name])
            )
            logger.info("Exported ungrouped report to: {}", file_path)

        return exported_files

    # ── Private helpers ───────────────────────────────────────────────────────

    def _write_workbook(
        self,
        file_path: Path,
        results: list[ReportResult],
    ) -> None:
        """Write one or more report results as sheets into a single workbook."""
        wb = Workbook()
        # Remove the default empty sheet created by openpyxl.
        wb.remove(wb.active)  # type: ignore[arg-type]

        for result in results:
            sheet_name = _sanitize_sheet_name(result.report_name)
            ws = wb.create_sheet(title=sheet_name)

            if result.status == "failed" or result.dataframe is None:
                ws.cell(row=1, column=1, value="Report failed — see log for details")
            else:
                self._write_dataframe(ws, result.dataframe)

        wb.save(file_path)

    @staticmethod
    def _write_dataframe(ws: object, df: pd.DataFrame) -> None:
        """Write a DataFrame to an openpyxl worksheet with header styling."""
        from openpyxl.worksheet.worksheet import Worksheet

        ws_typed: Worksheet = ws  # type: ignore[assignment]

        if df.empty:
            ws_typed.cell(row=1, column=1, value="(no data)")
            return

        columns = list(df.columns)

        # ── Write and style header row ────────────────────────────────────────
        for col_idx, col_name in enumerate(columns, start=1):
            cell = ws_typed.cell(row=1, column=col_idx, value=str(col_name))
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center")

        # ── Write data rows ───────────────────────────────────────────────────
        for row_idx, row_tuple in enumerate(df.itertuples(index=False), start=2):
            for col_idx, value in enumerate(row_tuple, start=1):
                ws_typed.cell(row=row_idx, column=col_idx, value=value)

        # ── Auto-fit column widths ────────────────────────────────────────────
        for col_idx, col_name in enumerate(columns, start=1):
            col_values = [str(col_name)] + [
                str(v) for v in df.iloc[:, col_idx - 1]
            ]
            max_width = max(len(v) for v in col_values) + 2  # +2 for padding
            ws_typed.column_dimensions[get_column_letter(col_idx)].width = min(
                max_width, 60  # cap at 60 to avoid huge columns
            )

        # ── Freeze header row ─────────────────────────────────────────────────
        ws_typed.freeze_panes = "A2"

        # ── Auto-filter on header row ─────────────────────────────────────────
        ws_typed.auto_filter.ref = ws_typed.dimensions

    @staticmethod
    def _resolve_path(out_dir: Path, filename_template: str, run_date: date) -> Path:
        """Replace {date} placeholder in filename with YYYY-MM-DD."""
        filename = filename_template.replace(
            "{date}", run_date.strftime("%Y-%m-%d")
        )
        return out_dir / filename


def _sanitize_sheet_name(name: str) -> str:
    """
    Make a string safe for use as an Excel sheet name:
      - Replace illegal characters (: \\ / ? * [ ]) with '_'
      - Truncate to 31 characters
    """
    safe = _ILLEGAL_SHEET_CHARS.sub("_", name)
    return safe[:_MAX_SHEET_NAME_LENGTH]
